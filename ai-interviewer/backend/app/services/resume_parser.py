"""Resume ingestion: file -> raw text -> structured ``resume_parsed`` dict.

The endpoint ``POST /api/v1/interview/resume/parse`` is intentionally a
*pre-processor* in front of the existing ``startSession`` flow: it
returns a JSON shape that the frontend can drop straight into the
``SetupForm`` first step, so the user always gets a chance to review
and correct the extraction before kicking off the LangGraph workflow.

Pipeline
--------

1. **Decode** the upload by content-type / extension into UTF-8 text.
   - ``.pdf`` via :mod:`pypdf` (already a transitive dep)
   - ``.docx`` via :mod:`docx` (``python-docx``)
   - ``.txt`` / ``.md`` decoded as UTF-8 with ``errors="replace"``
2. **Heuristic baseline** (always runs): cheap regex to pluck a
   summary, a skills list against a domain dictionary, and bulleted
   highlights. Guarantees a usable response even if the LLM fails or
   we are in stub mode.
3. **LLM refinement** (skipped when ``settings.use_stub_llm`` is true):
   ask the configured chat model to emit a tighter structured JSON blob
   (summary, skills, highlights, projects, focus areas, concerns). Only
   **non-empty** fields override the heuristic so a partial LLM failure
   can't blank out the heuristic output.

Both the truncated raw text and the final structured dict are returned
to the API layer; the raw text is handy for debugging / display but
NOT persisted by this module.
"""
from __future__ import annotations

import concurrent.futures
import contextvars
import hashlib
import io
import json
import re
import time
import unicodedata
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date
from typing import Any

from app.core.logging import get_logger
from app.core.settings import get_settings

log = get_logger(__name__)


MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_TEXT_CHARS = 60_000  # cap the LLM prompt size
MAX_PDF_PAGES = 25
MAX_DOCX_ZIP_ENTRIES = 256
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_ENTRY_BYTES = 8 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100.0
EXTRACT_TEXT_TIMEOUT_SECONDS = 10.0
LLM_REFINE_MAX_CHARS = 20_000
LLM_REFINE_MAX_OUTPUT_TOKENS = 4800
LLM_REFINE_REQUEST_TIMEOUT_SECONDS = 12.0
_EXTRACT_TEXT_EXECUTOR = concurrent.futures.ThreadPoolExecutor(
    max_workers=2,
    thread_name_prefix="resume-parser-extract",
)
_LLM_REFINE_EXECUTOR = concurrent.futures.ThreadPoolExecutor(
    max_workers=4,
    thread_name_prefix="resume-parser-llm",
)


class ResumeParseError(ValueError):
    """Raised when an upload cannot be decoded into usable text."""


def _parse_status(
    *,
    mode: str,
    reason: str,
    elapsed_ms: int | None = None,
    text_chars: int | None = None,
) -> dict[str, Any]:
    messages = {
        "ai_completed": "AI 已完成简历整理，下面内容可以直接检查和微调。",
        "heuristic": "已完成基础解析，可继续检查和微调。",
        "stub_mode": "已完成快速基础解析，可继续检查和微调。",
        "timeout": "已先完成基础整理，可继续检查和微调。",
        "llm_failed": "已先完成基础整理，智能补全暂时未完成。",
        "disabled": "已完成基础解析；AI 精修当前未启用。",
    }
    status: dict[str, Any] = {
        "mode": mode,
        "reason": reason,
        "message": messages.get(reason, messages["heuristic"]),
    }
    if elapsed_ms is not None:
        status["elapsed_ms"] = elapsed_ms
    if text_chars is not None:
        status["text_chars"] = text_chars
    return status


def _basic_parse_status() -> dict[str, Any]:
    return _parse_status(mode="basic", reason="heuristic")


@dataclass(frozen=True)
class ParsedResume:
    summary: str
    skills: list[str]
    highlights: list[str]
    raw_text: str
    candidate_name: str = ""
    candidate_profile: dict[str, Any] = field(default_factory=dict)
    projects: list[dict[str, Any]] = field(default_factory=list)
    focus_areas: list[dict[str, Any]] = field(default_factory=list)
    concerns: list[str] = field(default_factory=list)
    parse_status: dict[str, Any] = field(default_factory=_basic_parse_status)

    def to_dict(self) -> dict[str, Any]:
        return {
            "candidate_name": self.candidate_name,
            "candidate_profile": dict(self.candidate_profile),
            "summary": self.summary,
            "skills": list(self.skills),
            "highlights": list(self.highlights),
            "projects": list(self.projects),
            "focus_areas": list(self.focus_areas),
            "concerns": list(self.concerns),
            "raw_text": self.raw_text,
            "parse_status": dict(self.parse_status),
        }


# ---------------------------------------------------------------------------
# 1. File -> text
# ---------------------------------------------------------------------------


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ResumeParseError("pypdf is required for PDF resume parsing") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise ResumeParseError(
                f"PDF has too many pages ({page_count}); max {MAX_PDF_PAGES}"
            )
        chunks: list[str] = []
        for page in reader.pages:
            txt = page.extract_text() or ""
            if txt.strip():
                chunks.append(txt)
            if len("\n\n".join(chunks)) >= MAX_TEXT_CHARS:
                break
        return "\n\n".join(chunks)
    except ResumeParseError:
        raise
    except Exception as e:
        raise ResumeParseError(f"Failed to read PDF: {e}") from e


def _validate_docx_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as e:
        raise ResumeParseError("Invalid DOCX: not a valid zip container") from e

    names = {info.filename for info in infos}
    required = {"[Content_Types].xml", "word/document.xml"}
    if not required.issubset(names):
        missing = ", ".join(sorted(required - names))
        raise ResumeParseError(f"Invalid DOCX: missing required member(s): {missing}")

    if len(infos) > MAX_DOCX_ZIP_ENTRIES:
        raise ResumeParseError(
            f"DOCX has too many zip entries ({len(infos)}); max {MAX_DOCX_ZIP_ENTRIES}"
        )

    total_uncompressed = 0
    total_compressed = 0
    for info in infos:
        if info.flag_bits & 0x1:
            raise ResumeParseError(f"DOCX contains encrypted zip member: {info.filename}")
        total_uncompressed += int(info.file_size)
        total_compressed += int(info.compress_size)
        if info.file_size > MAX_DOCX_ENTRY_BYTES:
            raise ResumeParseError(
                "DOCX entry too large: "
                f"{info.filename} ({info.file_size} bytes); max {MAX_DOCX_ENTRY_BYTES}"
            )

    if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise ResumeParseError(
            "DOCX uncompressed size too large "
            f"({total_uncompressed} bytes); max {MAX_DOCX_UNCOMPRESSED_BYTES}"
        )

    ratio = total_uncompressed / max(total_compressed, 1)
    if ratio > MAX_DOCX_COMPRESSION_RATIO:
        raise ResumeParseError(
            "DOCX compression ratio too high "
            f"({ratio:.1f}x); max {MAX_DOCX_COMPRESSION_RATIO:.1f}x"
        )


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # type: ignore[import-untyped]
    except ImportError as e:  # pragma: no cover
        raise ResumeParseError("python-docx is required for DOCX resume parsing") from e
    try:
        _validate_docx_zip(data)
        document = docx.Document(io.BytesIO(data))
        parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        raise ResumeParseError(f"Failed to read DOCX: {e}") from e


def _extract_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"


def _declared_format(name: str, ctype: str) -> str | None:
    """Map filename/content_type to a declared format token.

    ``name``/``ctype`` are pre-lowered. ``None`` means the upload is in
    none of our supported families and should surface as
    ``unsupported`` upstream.
    """
    if name.endswith(".pdf") or "pdf" in ctype:
        return "pdf"
    if name.endswith(".docx") or "officedocument.wordprocessingml" in ctype:
        return "docx"
    if name.endswith((".txt", ".md", ".markdown")) or ctype.startswith("text/"):
        return "text"
    return None


def _verify_magic_bytes(data: bytes, declared: str) -> None:
    """Reject uploads whose declared format does not match the file header.

    Filename and content_type are user-controlled; without a magic-byte
    check, a malicious client could ship arbitrary bytes labelled as
    ``resume.pdf`` and force the parser deep into pypdf before failing.
    Mismatch surfaces as ``unsupported`` so the API maps it to 415.
    """
    if declared == "pdf" and not data.startswith(_PDF_MAGIC):
        raise ResumeParseError(
            "Unsupported file: declared as PDF but the file does not start "
            "with the PDF magic bytes (%PDF). Please upload the original PDF."
        )
    if declared == "docx" and not data.startswith(_ZIP_MAGIC):
        raise ResumeParseError(
            "Unsupported file: declared as DOCX but the file is not a valid "
            "ZIP container. Please upload the original .docx file."
        )


def extract_text(
    *,
    filename: str | None,
    content_type: str | None,
    data: bytes,
) -> str:
    """Decode an uploaded resume file into plain UTF-8 text.

    Raises :class:`ResumeParseError` when the file is too large, empty,
    or in an unsupported format. The caller should translate this into
    an HTTP 4xx.
    """
    if not data:
        raise ResumeParseError("Empty file uploaded")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeParseError(
            f"File too large ({len(data)} bytes); max {MAX_FILE_BYTES} bytes"
        )

    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    declared = _declared_format(name, ctype)
    if declared is None:
        raise ResumeParseError(
            f"Unsupported file type (filename={filename!r}, content_type={content_type!r}). "
            "Supported: .pdf .docx .txt .md"
        )
    _verify_magic_bytes(data, declared)

    if declared == "pdf":
        text = _extract_pdf(data)
    elif declared == "docx":
        text = _extract_docx(data)
    else:
        text = _extract_text(data)

    text = _clean_extracted_text(text).strip()
    if not text:
        raise ResumeParseError(
            "Could not extract any text from the file. "
            "If this is a scanned PDF, try uploading the source DOCX or pasting text."
        )
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    return text


def extract_text_with_timeout(
    *,
    filename: str | None,
    content_type: str | None,
    data: bytes,
    timeout_seconds: float = EXTRACT_TEXT_TIMEOUT_SECONDS,
) -> str:
    timeout = max(0.0, float(timeout_seconds))
    future = _EXTRACT_TEXT_EXECUTOR.submit(
        extract_text,
        filename=filename,
        content_type=content_type,
        data=data,
    )
    try:
        return future.result(timeout=timeout)
    except concurrent.futures.TimeoutError as e:
        future.cancel()
        raise ResumeParseError(
            f"Resume text extraction timed out after {timeout:.2f}s"
        ) from e


def _clean_extracted_text(text: str) -> str:
    """Normalize PDF/DOCX extraction artifacts before parsing.

    Some resume PDFs extract CJK compatibility glyphs (e.g. ``⽼``)
    and NUL/private-use markers around bullets. Normalizing here keeps
    both the heuristic parser and the LLM prompt closer to what a user
    actually sees on the page.
    """
    text = unicodedata.normalize("NFKC", text)
    cleaned_chars: list[str] = []
    for ch in text:
        if ch in "\n\r\t":
            cleaned_chars.append(ch)
            continue
        if ord(ch) < 32:
            continue
        if "\ue000" <= ch <= "\uf8ff":
            continue
        cleaned_chars.append(ch)
    return "".join(cleaned_chars)


# ---------------------------------------------------------------------------
# 2. Heuristic baseline
# ---------------------------------------------------------------------------

# Curated dictionary covering the skills our default rubrics + RAG corpora
# care about. Lower-case, hyphenated tokens. Adding here is cheap; missing
# entries are tolerated because the LLM step usually fills in the rest.
_SKILL_DICTIONARY = (
    # languages
    "python", "java", "javascript", "typescript", "go", "golang", "rust",
    "c++", "c#", "ruby", "php", "scala", "kotlin", "swift",
    # web frameworks
    "react", "vue", "next.js", "nextjs", "node.js", "nodejs", "express",
    "fastapi", "flask", "django", "spring", "spring-boot", "rails",
    # data & infra
    "postgres", "postgresql", "mysql", "mongodb", "redis", "kafka",
    "rabbitmq", "elasticsearch", "clickhouse", "snowflake", "spark",
    "airflow", "hadoop", "kubernetes", "k8s", "docker", "terraform",
    "aws", "gcp", "azure", "linux",
    # ml / ai
    "pytorch", "tensorflow", "langchain", "langgraph", "rag",
    "llm", "openai", "anthropic", "vector-db", "embeddings",
    # general
    "system-design", "system_design", "microservices", "event-driven",
    "ci/cd", "grpc", "rest", "graphql", "websocket",
)

_SKILL_ALIASES = (
    ("springboot", "spring"),
    ("spring boot", "spring"),
    ("springcloud", "spring"),
    ("spring cloud", "spring"),
    ("springsecurity", "spring"),
    ("spring security", "spring"),
    ("springai", "spring"),
    ("spring ai", "spring"),
)


def _contains_cjk(text: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


def _clip(text: str, limit: int = 80) -> str:
    text = re.sub(r"\s+", " ", text.strip())
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "…"


def _zh_join(items: list[str], *, limit: int = 5) -> str:
    visible = [item for item in items if item][:limit]
    if not visible:
        return ""
    if len(visible) == 1:
        return visible[0]
    return "、".join(visible)


def _heuristic_summary_raw(text: str) -> str:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    candidates = paragraphs[:3] if paragraphs else [text[:400]]
    summary = " ".join(candidates)
    summary = re.sub(r"\s+", " ", summary).strip()
    if len(summary) > 600:
        summary = summary[:600].rsplit(" ", 1)[0] + "…"
    return summary


def _heuristic_summary(
    text: str,
    *,
    skills: list[str],
    projects: list[dict[str, Any]],
    highlights: list[str],
) -> str:
    """Produce a Chinese fallback summary even when the source resume is English."""
    project_names = [
        str(project.get("name") or "").strip()
        for project in projects
        if str(project.get("name") or "").strip()
    ]
    anchors: list[str] = []
    for project in projects:
        anchors.extend(_normalise_list(project.get("question_anchors"), limit=3))
    anchors = list(dict.fromkeys(anchors))

    parts: list[str] = ["候选人简历已解析"]
    if project_names:
        parts.append(f"包含 {_zh_join([_clip(name, 28) for name in project_names], limit=3)} 等项目经历")
    elif highlights:
        parts.append("包含若干可追问的项目成果")
    if skills:
        parts.append(f"主要涉及 {_zh_join(skills, limit=8)} 等技术")
    if anchors:
        parts.append(f"面试可重点围绕 {_zh_join(anchors, limit=4)} 展开")

    if len(parts) == 1:
        raw = _heuristic_summary_raw(text)
        if _contains_cjk(raw):
            return raw
        return "候选人简历已提取到基础经历信息，建议补充项目、技能和成果，便于 AI 更准确地出题。"
    return "，".join(parts) + "。"


def _heuristic_skills(text: str) -> list[str]:
    lower = text.lower()
    found: list[str] = []
    seen: set[str] = set()
    for skill in _SKILL_DICTIONARY:
        pattern = r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"
        if re.search(pattern, lower) and skill not in seen:
            found.append(skill)
            seen.add(skill)
    for alias, canonical in _SKILL_ALIASES:
        if alias in lower and canonical not in seen:
            found.append(canonical)
            seen.add(canonical)
    return found


_BULLET_RE = re.compile(r"^\s*(?:[-*•·●▪◦]|\.(?=\s)|\d+[\.\)])\s+(.{20,})$")


def _heuristic_highlights(text: str) -> list[str]:
    out: list[str] = []
    for line in text.splitlines():
        m = _BULLET_RE.match(line)
        if not m:
            continue
        item = _friendly_highlight(re.sub(r"\s+", " ", m.group(1).strip()))
        if item and item not in out:
            out.append(item)
        if len(out) >= 6:
            break
    return out


def _friendly_highlight(item: str) -> str:
    if _contains_cjk(item):
        return item

    lower = item.lower()
    metric_match = re.search(
        r"(p\d{2,3}|latency|throughput|qps|events/day|teams?|ms|%)",
        item,
        re.IGNORECASE,
    )
    numbers = re.findall(r"\b\d+(?:\.\d+)?\s*(?:ms|s|%|m|k|w|teams?|events/day)?\b", item, re.IGNORECASE)
    metric_text = ""
    if metric_match or numbers:
        metric_items = [n.strip() for n in numbers]
        if metric_match:
            metric_items.insert(0, metric_match.group(0))
        metric_items = list(dict.fromkeys(metric_items))
        metric_text = f"（涉及指标：{_zh_join(metric_items, limit=3)}）" if metric_items else ""

    if any(token in lower for token in ("latency", "p99", "performance", "throughput", "reduced", "improved")):
        return f"性能优化：简历提到性能或延迟优化成果{metric_text}。"
    if any(token in lower for token in ("migration", "monolithic", "microservice", "event-driven", "architecture")):
        return "架构演进：简历提到系统迁移、微服务或事件驱动架构相关经历。"
    if any(token in lower for token in ("schema", "registry", "adopted", "teams", "company-wide")):
        return f"工程协作：简历提到跨团队复用或工程规范建设成果{metric_text}。"
    if any(token in lower for token in ("kafka", "pipeline", "events", "notification", "message")):
        return f"消息系统：简历提到消息管道或事件处理相关经历{metric_text}。"
    if any(token in lower for token in ("led", "owned", "designed", "built", "authored", "shipped")):
        return f"项目成果：简历提到主导、设计或交付相关成果{metric_text}。"
    return f"项目经历：{_clip(item, 60)}"


_SECTION_HEADINGS = {
    "summary",
    "profile",
    "skills",
    "technical skills",
    "experience",
    "work experience",
    "project experience",
    "projects",
    "education",
    "certifications",
}


def _normalise_list(items: Any, *, limit: int = 8) -> list[str]:
    if not isinstance(items, list):
        return []
    out: list[str] = []
    for item in items:
        if not isinstance(item, str):
            continue
        cleaned = re.sub(r"\s+", " ", item.strip())
        if cleaned and cleaned not in out:
            out.append(cleaned)
        if len(out) >= limit:
            break
    return out


_NAME_LABEL_RE = re.compile(r"^(?:name|candidate|姓名|候选人|称呼)\s*[:：]\s*", re.IGNORECASE)
_NAME_REJECT_TERMS = (
    "resume",
    "curriculum",
    "cv",
    "engineer",
    "developer",
    "architect",
    "manager",
    "summary",
    "profile",
    "skills",
    "experience",
    "education",
    "project",
    "java",
    "python",
    "backend",
    "frontend",
    "fullstack",
    "简历",
    "求职",
    "工程师",
    "开发",
    "后端",
    "前端",
    "全栈",
    "架构",
    "项目",
    "经历",
    "经验",
    "技能",
    "教育",
    "电话",
    "邮箱",
)


def _looks_like_candidate_name(value: str) -> bool:
    if not value or len(value) > 64:
        return False
    if re.search(r"@|https?://|\b1[3-9]\d{9}\b|\d{4,}", value, re.IGNORECASE):
        return False
    lower = value.lower()
    if any(term in lower for term in _NAME_REJECT_TERMS):
        return False

    cjk_chars = re.findall(r"[\u4e00-\u9fff]", value)
    if cjk_chars:
        # Common Chinese names are short; allow a small margin for compound surnames.
        return 2 <= len(cjk_chars) <= 6 and len(value.replace(" ", "")) <= 12

    words = re.findall(r"[A-Za-z][A-Za-z.'-]*", value)
    if not words or len(words) > 4:
        return False
    return len(" ".join(words)) >= 2


def _normalise_candidate_name(raw: Any) -> str:
    if not isinstance(raw, str):
        return ""
    value = re.sub(r"\s+", " ", raw.strip())
    value = value.lstrip("#*-•·").strip(" \t:-|")
    value = re.sub(r"\.(?:pdf|docx?|txt|md|markdown)$", "", value, flags=re.IGNORECASE)
    value = _NAME_LABEL_RE.sub("", value).strip()
    if not value:
        return ""

    for sep in (" | ", " - ", " — ", " – ", "\t"):
        if sep in value:
            for part in value.split(sep):
                name = _normalise_candidate_name(part)
                if name:
                    return name
            return ""

    return value if _looks_like_candidate_name(value) else ""


def _heuristic_candidate_name(text: str, filename: str | None = None) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines[:16]:
        labelled = _NAME_LABEL_RE.match(line)
        if labelled:
            name = _normalise_candidate_name(line)
            if name:
                return name
    for line in lines[:8]:
        name = _normalise_candidate_name(line)
        if name:
            return name
    return _normalise_candidate_name(filename or "")


_EDUCATION_LEVELS = ("博士", "硕士", "研究生", "本科", "大专", "专科", "高中")
_ROLE_TERMS = (
    "engineer",
    "developer",
    "architect",
    "manager",
    "java",
    "backend",
    "frontend",
    "fullstack",
    "后端",
    "前端",
    "全栈",
    "架构",
    "开发",
    "工程师",
    "算法",
    "运维",
    "测试",
)

_JOB_LEVELS = {"junior", "mid", "senior", "staff", "principal"}


def _clean_profile_text(raw: Any, *, limit: int = 80) -> str:
    if not isinstance(raw, str):
        return ""
    value = re.sub(r"\s+", " ", raw.strip())
    value = value.strip(" \t:-|,，。；;")
    if not value or len(value) > limit:
        return ""
    if re.search(r"@|https?://|\b1[3-9]\d{9}\b", value, re.IGNORECASE):
        return ""
    return value


def _normalise_major(raw: Any) -> str:
    value = _clean_profile_text(raw, limit=60)
    value = re.sub(r"(?:专业|major)", "", value, flags=re.IGNORECASE)
    return re.sub(r"^[\s:\-|,，。；;.]+|[\s:\-|,，。；;.]+$", "", value)


def _normalise_role(raw: Any) -> str:
    value = _clean_profile_text(raw, limit=80)
    if not value:
        return ""
    value = re.sub(r"\bwith\s+\d+(?:\.\d+)?\+?\s+years?.*$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"(?:\d+(?:\.\d+)?\+?\s*年.*)$", "", value).strip()
    value = re.sub(r"([A-Za-z])(?=[\u4e00-\u9fff])", r"\1 ", value)
    value = re.sub(r"(?<=[\u4e00-\u9fff])([A-Za-z])", r" \1", value)
    value = re.sub(r"\s+", " ", value).strip(" -–—|")
    lower = value.lower()
    if not value or not any(term in lower for term in _ROLE_TERMS):
        return ""
    return value


def _normalise_job_title(raw: Any) -> str:
    role = _normalise_role(raw)
    if not role:
        return ""
    compact = role.replace(" ", "")
    if "Java后端" in compact and "工程师" not in role:
        return "Java 后端开发工程师"
    if "后端" in role and "工程师" not in role:
        return f"{role}开发工程师"
    if re.search(r"\bbackend\b", role, re.IGNORECASE) and "Engineer" not in role:
        return f"{role} Engineer"
    return role


def _normalise_experience_years(raw: Any) -> int | float | None:
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        years = float(raw)
    elif isinstance(raw, str):
        m = re.search(r"(\d+(?:\.\d+)?)", raw)
        if not m:
            return None
        years = float(m.group(1))
    else:
        return None
    if years < 0 or years > 60:
        return None
    return int(years) if years.is_integer() else round(years, 1)


def _normalise_year(raw: Any, *, min_year: int = 1950, max_year: int | None = None) -> int | None:
    if max_year is None:
        max_year = date.today().year + 10
    if isinstance(raw, int) and not isinstance(raw, bool):
        year = raw
    elif isinstance(raw, str):
        m = re.search(r"\b(19\d{2}|20\d{2})\b", raw)
        if not m:
            return None
        year = int(m.group(1))
    else:
        return None
    return year if min_year <= year <= max_year else None


def _normalise_age(raw: Any) -> int | None:
    if isinstance(raw, int) and not isinstance(raw, bool):
        age = raw
    elif isinstance(raw, str):
        m = re.search(r"\b(\d{1,2})\b", raw)
        if not m:
            return None
        age = int(m.group(1))
    else:
        return None
    return age if 14 <= age <= 80 else None


def _normalise_bool(raw: Any) -> bool | None:
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, str):
        value = raw.strip().lower()
        if value in {"true", "yes", "y", "1", "是", "应届", "应届生"}:
            return True
        if value in {"false", "no", "n", "0", "否", "非应届"}:
            return False
    return None


def _normalise_basis(raw: Any) -> list[str]:
    if not isinstance(raw, list):
        return []
    out: list[str] = []
    for item in raw:
        value = _clean_profile_text(item, limit=80)
        if value and value not in out:
            out.append(value)
        if len(out) >= 4:
            break
    return out


def _normalise_candidate_profile(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        return {}
    out: dict[str, Any] = {}
    education_level = _clean_profile_text(raw.get("education_level"), limit=20)
    if education_level:
        out["education_level"] = education_level
    school = _clean_profile_text(raw.get("school"), limit=80)
    if school:
        out["school"] = school
    major = _normalise_major(raw.get("major"))
    if major:
        out["major"] = major
    years = _normalise_experience_years(raw.get("experience_years"))
    if years is not None:
        out["experience_years"] = years
    birth_year = _normalise_year(raw.get("birth_year"), min_year=1940, max_year=date.today().year)
    if birth_year is not None:
        out["birth_year"] = birth_year
    age = _normalise_age(raw.get("age"))
    if age is None and birth_year is not None:
        age = max(0, date.today().year - birth_year)
    if age is not None:
        out["age"] = age
    graduation_year = _normalise_year(raw.get("graduation_year"), min_year=1950)
    if graduation_year is not None:
        out["graduation_year"] = graduation_year
    fresh_graduate = _normalise_bool(raw.get("fresh_graduate"))
    if fresh_graduate is not None:
        out["fresh_graduate"] = fresh_graduate
    role = _normalise_role(raw.get("current_or_target_role"))
    if role:
        out["current_or_target_role"] = role
    suggested_title = _normalise_job_title(raw.get("suggested_job_title"))
    if suggested_title:
        out["suggested_job_title"] = suggested_title
    level = raw.get("suggested_job_level")
    if isinstance(level, str) and level in _JOB_LEVELS:
        out["suggested_job_level"] = level
    basis = _normalise_basis(raw.get("suggested_job_level_basis"))
    if basis:
        out["suggested_job_level_basis"] = basis
    return out


def _education_from_line(line: str) -> dict[str, Any]:
    profile: dict[str, Any] = {}
    if not any(level in line for level in _EDUCATION_LEVELS) and not re.search(
        r"university|college|b\.?s\.?|m\.?s\.?|ph\.?d", line, re.IGNORECASE
    ):
        return profile

    parts = [
        p.strip()
        for p in re.split(r"\s*(?:[-–—|/])\s*", line)
        if p.strip()
    ]
    for part in parts:
        if "大学" in part or "学院" in part or re.search(r"university|college", part, re.IGNORECASE):
            profile["school"] = _clean_profile_text(part)
            break
    for level in _EDUCATION_LEVELS:
        if level in line:
            profile["education_level"] = level
            break
    for part in parts:
        major = _normalise_major(part)
        if major and ("专业" in part or "major" in part.lower()):
            profile["major"] = major
            break
    if "major" not in profile:
        m = re.search(r"([\u4e00-\u9fffA-Za-z0-9+#.\s]{2,40})专业", line)
        if m:
            profile["major"] = _normalise_major(m.group(1))
    return _normalise_candidate_profile(profile)


def _experience_years_from_text(text: str) -> int | float | None:
    patterns = (
        r"(\d+(?:\.\d+)?)\+?\s*年(?:以上)?(?:[^。\n]{0,12})?(?:经验|经历|开发)",
        r"(\d+(?:\.\d+)?)\+?\s*years?(?:\s+of)?(?:[^.\n]{0,30})?(?:experience|building|developing)",
    )
    for pattern in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            return _normalise_experience_years(m.group(1))
    return None


def _birth_year_from_text(text: str) -> int | None:
    patterns = (
        r"(?:男|女)\s*/\s*(19\d{2}|20\d{2})",
        r"(19\d{2}|20\d{2})\s*年\s*(?:出生|生)",
    )
    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return _normalise_year(m.group(1), min_year=1940, max_year=date.today().year)
    return None


def _graduation_year_from_text(text: str) -> int | None:
    patterns = (
        r"(20\d{2})\s*届",
        r"(20\d{2})\s*年\s*(?:毕业|毕业生)",
        r"(?:毕业|graduat(?:e|ion))[^\n]{0,12}(20\d{2})",
    )
    for pattern in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            return _normalise_year(m.group(1), min_year=1950)
    return None


def _fresh_graduate_from_text(text: str, graduation_year: int | None) -> bool | None:
    if re.search(r"应届|在校|校招|实习生|new grad|fresh graduate", text, re.IGNORECASE):
        return True
    if re.search(r"社招|全职工作|工作经验", text):
        return False
    if graduation_year is not None:
        current_year = date.today().year
        if graduation_year - 1 <= current_year <= graduation_year + 1:
            return True
    return None


def _suggest_job_level(profile: dict[str, Any]) -> tuple[str | None, list[str]]:
    years = _normalise_experience_years(profile.get("experience_years"))
    if years is not None:
        label = f"{years:g} 年全职/工作经验"
        if years >= 8:
            return "staff", [label]
        if years >= 5:
            return "senior", [label]
        if years >= 2:
            return "mid", [label]
        return "junior", [label]

    basis: list[str] = []
    if profile.get("fresh_graduate") is True:
        basis.append("应届/在读")
    if profile.get("education_level"):
        basis.append(str(profile["education_level"]))
    if profile.get("graduation_year"):
        basis.append(f"{profile['graduation_year']} 届/毕业")
    if profile.get("age"):
        basis.append(f"{profile['age']} 岁")
    if basis:
        basis.append("未发现明确全职工作经验")
        return "junior", basis
    return None, []


def _with_profile_suggestions(profile: dict[str, Any]) -> dict[str, Any]:
    out = _normalise_candidate_profile(profile)
    title = _normalise_job_title(out.get("current_or_target_role")) or _normalise_job_title(
        out.get("suggested_job_title")
    )
    if title:
        out["suggested_job_title"] = title
    level, basis = _suggest_job_level(out)
    if level:
        out["suggested_job_level"] = level
        out["suggested_job_level_basis"] = basis
    return out


def _role_from_lines(lines: list[str], candidate_name: str) -> str:
    for line in lines[:8]:
        if "@" in line or re.search(r"\b1[3-9]\d{9}\b", line):
            continue
        parts = [
            p.strip()
            for p in re.split(r"\s*(?:[-–—|/])\s*", line)
            if p.strip()
        ]
        for part in parts:
            if candidate_name and part == candidate_name:
                continue
            role = _normalise_role(part)
            if role:
                return role
        role = _normalise_role(line)
        if role and role != candidate_name:
            return role
    return ""


def _heuristic_candidate_profile(text: str, filename: str | None = None) -> dict[str, Any]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    head_text = "\n".join(lines[:40])
    profile: dict[str, Any] = {}
    for line in lines[:40]:
        profile.update(_education_from_line(line))

    years = _experience_years_from_text(head_text)
    if years is not None:
        profile["experience_years"] = years
    birth_year = _birth_year_from_text(head_text)
    if birth_year is not None:
        profile["birth_year"] = birth_year
    graduation_year = _graduation_year_from_text(head_text)
    if graduation_year is not None:
        profile["graduation_year"] = graduation_year
    fresh_graduate = _fresh_graduate_from_text(head_text, graduation_year)
    if fresh_graduate is not None:
        profile["fresh_graduate"] = fresh_graduate

    candidate_name = _heuristic_candidate_name(text, filename)
    role = _role_from_lines(lines, candidate_name)
    if not role and filename:
        role = _role_from_lines([filename], candidate_name)
    if role:
        profile["current_or_target_role"] = role

    return _with_profile_suggestions(profile)


def _skills_in_text(text: str) -> list[str]:
    lower = text.lower()
    found: list[str] = []
    for skill in _SKILL_DICTIONARY:
        pattern = r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    for alias, canonical in _SKILL_ALIASES:
        if alias in lower and canonical not in found:
            found.append(canonical)
    return found


def _looks_like_project_header(line: str) -> bool:
    stripped = line.strip()
    if not stripped or _BULLET_RE.match(stripped):
        return False
    lower = stripped.lower().strip(":：")
    if lower in _SECTION_HEADINGS:
        return False
    non_project_prefixes = (
        "专业技能",
        "框架技术",
        "数据库",
        "微服务",
        "中间件",
        "开发工具",
        "编程语言",
        "后端技术",
        "前端技术",
        "ai技术",
        "ai 技术",
        "其他技能",
    )
    if lower.startswith(non_project_prefixes):
        return False
    if re.search(r"@|1[3-9]\d{9}|男/|女/|电话|邮箱|email", stripped, re.IGNORECASE):
        return False
    if len(stripped) < 8 or len(stripped) > 140:
        return False
    has_year = bool(re.search(r"(?:19|20)\d{2}", stripped))
    has_separator = any(sep in stripped for sep in (" | ", " - ", " — ", " – "))
    project_words = (
        "project",
        "platform",
        "system",
        "service",
        "pipeline",
        "migration",
        "architecture",
        "console",
        "app",
        "项目",
        "平台",
        "系统",
        "服务",
        "迁移",
        "架构",
    )
    return (has_year and has_separator) or any(w in lower for w in project_words)


def _project_name_and_role(header: str) -> tuple[str, str]:
    cleaned = re.sub(r"\([^)]*(?:19|20)\d{2}[^)]*\)", "", header).strip()
    cleaned = re.sub(
        r"\s*(?:19|20)\d{2}年\d{1,2}月\s*[-－–—~至到]+\s*(?:(?:19|20)\d{2}年\d{1,2}月|至今|present|now)\s*",
        " ",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"\b(?:19|20)\d{2}\s*[-–—~至到]\s*(?:present|now|至今|(?:19|20)\d{2})\b", "", cleaned, flags=re.IGNORECASE).strip(" -–—|")
    for sep in (" | ", " — ", " – ", " - "):
        if sep in cleaned:
            left, _, right = cleaned.partition(sep)
            return left.strip() or cleaned, right.strip()
    return cleaned, ""


def _question_anchors(text: str, tech_stack: list[str]) -> list[str]:
    lower = text.lower()
    anchors: list[str] = []
    rules = [
        (("migration", "architecture", "microservice", "distributed", "架构", "迁移"), "架构选型"),
        (("latency", "p99", "performance", "throughput", "性能", "延迟"), "性能优化"),
        (("kafka", "event", "queue", "idempot", "一致性", "消息"), "消息与一致性"),
        (("redis", "cache", "缓存"), "缓存设计"),
        (("kubernetes", "k8s", "docker", "deploy", "rollout", "发布"), "稳定性与发布"),
        (("model", "llm", "rag", "agent", "模型"), "AI 工程落地"),
    ]
    for keys, label in rules:
        if any(k in lower for k in keys) and label not in anchors:
            anchors.append(label)
    for skill in tech_stack[:3]:
        label = f"{skill} 项目应用"
        if label not in anchors:
            anchors.append(label)
    if not anchors:
        anchors = ["项目背景", "技术选型", "结果复盘"]
    return anchors[:5]


def _dimensions_for_anchor(label: str) -> list[str]:
    if any(k in label for k in ("架构", "一致性", "缓存", "稳定性", "发布")):
        return ["system_design", "technical_depth"]
    if any(k in label for k in ("性能", "优化", "结果", "复盘")):
        return ["problem_solving", "technical_depth"]
    if any(k in label for k in ("AI", "工程", "模型")):
        return ["technical_depth", "product_thinking"]
    return ["project_experience", "technical_depth"]


def _heuristic_projects(text: str) -> list[dict[str, Any]]:
    lines = [line.rstrip() for line in text.splitlines()]
    groups: list[tuple[str, list[str]]] = []
    current_header = ""
    current_bullets: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if _looks_like_project_header(line):
            if current_header:
                groups.append((current_header, current_bullets))
            current_header = line
            current_bullets = []
            continue
        m = _BULLET_RE.match(line)
        if m and current_header:
            current_bullets.append(re.sub(r"\s+", " ", m.group(1).strip()))
            continue
        if current_header and re.match(r"^(技术栈|技术亮点|项目职责|职责|成果)\s*[:：]?", line):
            current_bullets.append(re.sub(r"\s+", " ", line.strip()))
    if current_header:
        groups.append((current_header, current_bullets))

    projects: list[dict[str, Any]] = []
    for idx, (header, bullets) in enumerate(groups[:5], start=1):
        body = " ".join([header, *bullets])
        tech_stack = _skills_in_text(body)
        achievements = [
            b
            for b in bullets
            if re.search(r"\d|reduced|improved|led|designed|built|shipped|owned|优化|提升|负责|主导|设计", b, re.IGNORECASE)
        ] or bullets[:3]
        name, role = _project_name_and_role(header)
        if not name:
            continue
        projects.append(
            {
                "id": f"proj-{idx}",
                "name": name[:120],
                "role": role[:120],
                "tech_stack": tech_stack[:10],
                "responsibilities": bullets[:4],
                "achievements": achievements[:4],
                "question_anchors": _question_anchors(body, tech_stack),
            }
        )
    return projects


def _derive_focus_areas(
    *,
    projects: list[dict[str, Any]],
    skills: list[str],
    highlights: list[str],
) -> list[dict[str, Any]]:
    focus: list[dict[str, Any]] = []
    for project in projects[:4]:
        anchors = _normalise_list(project.get("question_anchors"), limit=3)
        for anchor in anchors[:2]:
            focus.append(
                {
                    "id": f"focus-{len(focus) + 1}",
                    "label": f"{project.get('name', '项目')}：{anchor}",
                    "project_id": project.get("id"),
                    "dimensions": _dimensions_for_anchor(anchor),
                    "skills": list(project.get("tech_stack") or [])[:5],
                    "priority": len(focus) + 1,
                }
            )
            if len(focus) >= 6:
                return focus
    if not focus:
        for highlight in highlights[:3]:
            focus.append(
                {
                    "id": f"focus-{len(focus) + 1}",
                    "label": highlight[:80],
                    "project_id": None,
                    "dimensions": ["project_experience", "technical_depth"],
                    "skills": skills[:5],
                    "priority": len(focus) + 1,
                }
            )
    if not focus:
        for skill in skills[:5]:
            focus.append(
                {
                    "id": f"focus-{len(focus) + 1}",
                    "label": f"{skill} 的项目应用与基础理解",
                    "project_id": None,
                    "dimensions": ["technical_depth"],
                    "skills": [skill],
                    "priority": len(focus) + 1,
                }
            )
    return focus[:6]


def _derive_concerns(projects: list[dict[str, Any]], highlights: list[str]) -> list[str]:
    concerns: list[str] = []
    for project in projects:
        if project.get("tech_stack") and not project.get("achievements"):
            concerns.append(
                f"{project.get('name', '项目')} 提到了技术栈，但成果或权衡细节较少。"
            )
        if not project.get("question_anchors"):
            concerns.append(
                f"{project.get('name', '项目')} 可以补充更多可追问的技术细节。"
            )
        if len(concerns) >= 3:
            break
    if not projects and highlights:
        concerns.append("简历亮点较分散，面试会先让你补充具体项目背景。")
    return concerns


def heuristic_parse(text: str, *, filename: str | None = None) -> ParsedResume:
    text = _clean_extracted_text(text)
    skills = _heuristic_skills(text)
    highlights = _heuristic_highlights(text)
    projects = _heuristic_projects(text)
    return ParsedResume(
        candidate_name=_heuristic_candidate_name(text, filename),
        candidate_profile=_heuristic_candidate_profile(text, filename),
        summary=_heuristic_summary(
            text,
            skills=skills,
            projects=projects,
            highlights=highlights,
        ),
        skills=skills,
        highlights=highlights,
        raw_text=text,
        projects=projects,
        focus_areas=_derive_focus_areas(
            projects=projects,
            skills=skills,
            highlights=highlights,
        ),
        concerns=_derive_concerns(projects, highlights),
    )


# ---------------------------------------------------------------------------
# 3. LLM refinement
# ---------------------------------------------------------------------------

_SYSTEM_PROMPT = """你是一名严谨的简历解析器。

请阅读用户提供的简历文本，并且只输出一个 JSON 对象，字段如下：

- "summary": 2-4 句简体中文，概括候选人的资历、主要领域和最强的亮点。
  只写给用户看的内容，不要营销话术。
- "candidate_name": 候选人的姓名或简短称呼。只在简历中明确出现时填写；
  不要把岗位名称、电话、邮箱、学校、公司或项目名当成姓名。
- "candidate_profile": 非敏感基本信息对象，只包含 education_level、school、
  major、experience_years、current_or_target_role、birth_year、age、
  graduation_year、fresh_graduate、suggested_job_title、suggested_job_level、
  suggested_job_level_basis。不要输出性别、婚育、民族、籍贯、头像、
  手机号、邮箱、身份证或地址。
  suggested_job_level 只能根据应届/在读、毕业年份、学历、明确全职工作年限、
  简历明示职级推断；不要根据项目技术栈复杂度提高级别。
- "skills": 归一化后的技能关键词数组，使用小写英文或常见技术写法，
  只保留简历中明确出现的技能（例如 "kubernetes"、"system-design"、
  "next.js"）。去重。
- "highlights": 3-6 条具体成果，使用简体中文单句表达。优先保留带有
  指标、规模或结果的内容（例如“将 p99 从 800ms 降到 120ms”、
  “推动 12 个团队共用的单体仓库迁移到事件驱动架构”）。
- "projects": 最多 5 个具体项目 / 经历锚点。每项包含 id、name、role、
  tech_stack、responsibilities、achievements、question_anchors。
  其中用户可见文本都使用简体中文；数组字段每项最多 3 条，每条保持短句。
- "focus_areas": 最多 6 个面试重点，来源于项目经历。每项包含 id、label、
  project_id、dimensions、skills、priority。label 使用简体中文。
- "concerns": 简短中文说明缺失或证据较弱、面试中适合追问的点。

硬性规则：
- 只输出合法 JSON。不要 Markdown 代码块，不要解释。
- 除了 "skills" 和 "dimensions" 以外，所有用户可见文本都必须使用简体中文。
- 输出必须是完整闭合的 JSON 对象，优先保证 JSON 有效，不要为了写满细节导致截断。
- 如果某个字段缺失或不清楚，就返回空值
  （summary "", skills [], highlights [], projects [], focus_areas [],
  concerns [], candidate_name "", candidate_profile {}），不要编造事实。
- 不要直接逐字引用原始简历中过长的句子；尽量改写成简短中文。
"""


def _clip_text_for_llm_refine(text: str) -> str:
    """Keep resume parsing responsive on long PDF/DOCX extracts."""
    if len(text) <= LLM_REFINE_MAX_CHARS:
        return text
    head_chars = int(LLM_REFINE_MAX_CHARS * 0.7)
    tail_chars = LLM_REFINE_MAX_CHARS - head_chars
    omitted = len(text) - head_chars - tail_chars
    return (
        text[:head_chars].rstrip()
        + "\n\n[content clipped: 为避免解析超时，中间省略 "
        + str(max(0, omitted))
        + " 个字符；请只根据可见内容整理，不要编造缺失信息。]\n\n"
        + text[-tail_chars:].lstrip()
    )


def _trace_resume_inputs(inputs: dict[str, Any]) -> dict[str, Any]:
    text = inputs.get("text")
    if not isinstance(text, str):
        return {"input": "non_text"}
    return {
        "text_chars": len(text),
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest()[:16],
        "llm_text_chars": len(_clip_text_for_llm_refine(text)),
    }


def _trace_resume_outputs(outputs: Any) -> dict[str, Any]:
    if not isinstance(outputs, dict):
        return {"output_type": type(outputs).__name__}
    return {
        "has_candidate_name": bool(outputs.get("candidate_name")),
        "has_candidate_profile": bool(outputs.get("candidate_profile")),
        "has_summary": bool(outputs.get("summary")),
        "skills_count": len(outputs.get("skills") or []),
        "projects_count": len(outputs.get("projects") or []),
        "focus_areas_count": len(outputs.get("focus_areas") or []),
        "concerns_count": len(outputs.get("concerns") or []),
    }


def _resume_traceable():
    try:
        from langsmith import traceable
    except Exception:  # pragma: no cover - optional observability dependency
        return lambda fn: fn
    return traceable(
        name="resume_parser_refine",
        run_type="chain",
        process_inputs=_trace_resume_inputs,
        process_outputs=_trace_resume_outputs,
    )


def _strip_json_fence(text: str) -> str:
    text = (text or "").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:]
    return text.strip()


def _repair_truncated_json_object(text: str) -> dict[str, Any]:
    """Best-effort repair for provider responses cut off near the end.

    Resume parsing is a pre-fill helper. When a provider returns a clear JSON
    prefix but misses the final brackets, keeping the usable fields is better
    than throwing away the whole AI pass.
    """
    candidate = _strip_json_fence(text)
    start = candidate.find("{")
    if start == -1:
        return {}
    candidate = candidate[start:]

    out: list[str] = []
    stack: list[str] = []
    in_string = False
    escaped = False

    for ch in candidate:
        if in_string:
            if escaped:
                out.append(ch)
                escaped = False
                continue
            if ch == "\\":
                out.append(ch)
                escaped = True
                continue
            if ch == '"':
                out.append(ch)
                in_string = False
                continue
            if ch in "\r\n":
                out.append("\\n")
                continue
            out.append(ch)
            continue

        if ch == '"':
            out.append(ch)
            in_string = True
        elif ch == "{":
            out.append(ch)
            stack.append("}")
        elif ch == "[":
            out.append(ch)
            stack.append("]")
        elif ch in "}]":
            if stack and stack[-1] == ch:
                stack.pop()
                out.append(ch)
                if not stack:
                    break
            else:
                break
        else:
            out.append(ch)

    if in_string:
        out.append('"')

    repaired = "".join(out).rstrip()
    while repaired.endswith((",", ":")):
        repaired = repaired[:-1].rstrip()
    while stack:
        repaired += stack.pop()
    repaired = re.sub(r",\s*([}\]])", r"\1", repaired)

    try:
        parsed = json.loads(repaired)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


@_resume_traceable()
def _llm_refine(
    text: str,
    *,
    request_timeout: float | None = None,
) -> dict[str, Any]:
    """Try to enrich the heuristic result via the configured LLM.

    Returns ``{}`` (silently) on any failure so the caller falls back
    to the heuristic baseline. Logs at warning level so operators can
    see provider-side issues without breaking the user's flow.
    """
    from app.engine.agents.llm_client import (
        ChatMessage,
        call_chat,
        parse_json_response,
    )

    llm_text = _clip_text_for_llm_refine(text)
    messages = [
        ChatMessage(role="system", content=_SYSTEM_PROMPT),
        ChatMessage(
            role="user",
            content=(
                "请把下面的简历整理成上面定义的 JSON 结构。\n"
                "除 skills 和 dimensions 外，所有用户可见文本都用简体中文。\n\n"
                "----- 简历开始 -----\n"
                f"{llm_text}\n"
                "----- 简历结束 -----"
            ),
        ),
    ]
    try:
        raw = call_chat(
            messages,
            json_mode=True,
            temperature=0.1,
            max_tokens=LLM_REFINE_MAX_OUTPUT_TOKENS,
            request_timeout=(
                LLM_REFINE_REQUEST_TIMEOUT_SECONDS
                if request_timeout is None
                else request_timeout
            ),
            max_retries=0,
            provider_max_retries=0,
            agent_role="resume_parser",
        )
    except Exception as e:
        try:
            from app.engine.agents.llm_client import redact_llm_secrets
            from app.services.session_manager import get_llm_override

            safe_error = redact_llm_secrets(str(e), get_llm_override())
        except Exception:
            safe_error = str(e)
        log.warning(
            "resume_parser: LLM call failed, falling back to heuristic: %s",
            safe_error,
        )
        return {}
    parsed = parse_json_response(raw)
    if not isinstance(parsed, dict) or not parsed:
        repaired = _repair_truncated_json_object(raw)
        if repaired:
            log.warning("resume_parser: repaired malformed LLM JSON response")
            parsed = repaired
    if not isinstance(parsed, dict):
        return {}
    return parsed


def _normalise_skill(s: Any) -> str | None:
    if not isinstance(s, str):
        return None
    cleaned = re.sub(r"\s+", "-", s.strip().lower())
    return cleaned or None


def _normalise_projects(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list):
        return []
    projects: list[dict[str, Any]] = []
    for idx, item in enumerate(raw[:5], start=1):
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        projects.append(
            {
                "id": str(item.get("id") or f"proj-{idx}")[:40],
                "name": name[:120],
                "role": str(item.get("role") or "").strip()[:120],
                "tech_stack": [
                    s.lower() for s in _normalise_list(item.get("tech_stack"), limit=10)
                ],
                "responsibilities": _normalise_list(
                    item.get("responsibilities"),
                    limit=5,
                ),
                "achievements": _normalise_list(item.get("achievements"), limit=5),
                "question_anchors": _normalise_list(
                    item.get("question_anchors"),
                    limit=5,
                ),
            }
        )
    return projects


def _normalise_focus_areas(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list):
        return []
    focus: list[dict[str, Any]] = []
    for idx, item in enumerate(raw[:6], start=1):
        if not isinstance(item, dict):
            continue
        label = str(item.get("label") or "").strip()
        if not label:
            continue
        try:
            priority = int(item.get("priority") or idx)
        except (TypeError, ValueError):
            priority = idx
        project_id = item.get("project_id")
        focus.append(
            {
                "id": str(item.get("id") or f"focus-{idx}")[:40],
                "label": label[:120],
                "project_id": str(project_id)[:40] if project_id else None,
                "dimensions": _normalise_list(item.get("dimensions"), limit=5),
                "skills": _normalise_list(item.get("skills"), limit=8),
                "priority": priority,
            }
        )
    return focus


def _merge(heuristic: ParsedResume, llm: dict[str, Any]) -> ParsedResume:
    """Layer LLM output over the heuristic baseline.

    Only non-empty LLM fields override; this keeps stub mode and
    failed calls from blanking out a perfectly good heuristic result.
    Skill lists from both sources are merged (LLM-first ordering).
    """
    candidate_name = heuristic.candidate_name
    llm_candidate_name = _normalise_candidate_name(llm.get("candidate_name"))
    if llm_candidate_name:
        candidate_name = llm_candidate_name

    candidate_profile = dict(heuristic.candidate_profile)
    llm_candidate_profile = _normalise_candidate_profile(llm.get("candidate_profile"))
    if llm_candidate_profile:
        candidate_profile.update(llm_candidate_profile)
    candidate_profile = _with_profile_suggestions(candidate_profile)

    summary = heuristic.summary
    if isinstance(llm.get("summary"), str) and llm["summary"].strip():
        summary = llm["summary"].strip()

    llm_skills_raw = llm.get("skills")
    merged_skills: list[str] = []
    seen: set[str] = set()
    if isinstance(llm_skills_raw, list):
        for s in llm_skills_raw:
            n = _normalise_skill(s)
            if n and n not in seen:
                merged_skills.append(n)
                seen.add(n)
    for s in heuristic.skills:
        if s not in seen:
            merged_skills.append(s)
            seen.add(s)

    highlights = heuristic.highlights
    llm_highlights = llm.get("highlights")
    if isinstance(llm_highlights, list):
        normalised = [
            re.sub(r"\s+", " ", h.strip())
            for h in llm_highlights
            if isinstance(h, str) and h.strip()
        ]
        if normalised:
            highlights = normalised[:6]

    projects = heuristic.projects
    llm_projects = _normalise_projects(llm.get("projects"))
    if llm_projects:
        projects = llm_projects

    focus_areas = heuristic.focus_areas
    llm_focus_areas = _normalise_focus_areas(llm.get("focus_areas"))
    if llm_focus_areas:
        focus_areas = llm_focus_areas
    elif projects != heuristic.projects or highlights != heuristic.highlights:
        focus_areas = _derive_focus_areas(
            projects=projects,
            skills=merged_skills,
            highlights=highlights,
        )

    concerns = heuristic.concerns
    llm_concerns = _normalise_list(llm.get("concerns"), limit=5)
    if llm_concerns:
        concerns = llm_concerns

    return ParsedResume(
        candidate_name=candidate_name,
        candidate_profile=candidate_profile,
        summary=summary,
        skills=merged_skills,
        highlights=highlights,
        raw_text=heuristic.raw_text,
        projects=projects,
        focus_areas=focus_areas,
        concerns=concerns,
        parse_status=heuristic.parse_status,
    )


def _with_parse_status(
    parsed: ParsedResume,
    *,
    mode: str,
    reason: str,
    started_at: float,
) -> ParsedResume:
    return replace(
        parsed,
        parse_status=_parse_status(
            mode=mode,
            reason=reason,
            elapsed_ms=max(0, int((time.perf_counter() - started_at) * 1000)),
            text_chars=len(parsed.raw_text),
        ),
    )


def parse_resume(
    text: str,
    *,
    force_llm: bool = False,
    filename: str | None = None,
) -> ParsedResume:
    """Public entrypoint: heuristic baseline + optional LLM refinement.

    Stub mode (``settings.use_stub_llm``) skips the LLM call entirely
    so that the demo loop with no API key still produces a clean
    extraction without spending tokens on the deterministic fixture.
    """
    started_at = time.perf_counter()
    baseline = heuristic_parse(text, filename=filename)
    if get_settings().use_stub_llm and not force_llm:
        return _with_parse_status(
            baseline,
            mode="basic",
            reason="stub_mode",
            started_at=started_at,
        )
    timeout = max(0.0, float(get_settings().resume_parser_llm_timeout_seconds))
    if timeout <= 0:
        log.warning(
            "resume_parser: LLM refinement disabled by timeout=%.2f; using heuristic",
            timeout,
        )
        return _with_parse_status(
            baseline,
            mode="basic",
            reason="disabled",
            started_at=started_at,
        )
    ctx = contextvars.copy_context()
    future = _LLM_REFINE_EXECUTOR.submit(
        ctx.run,
        _llm_refine,
        text,
        request_timeout=timeout,
    )
    try:
        refined = future.result(timeout=timeout)
    except concurrent.futures.TimeoutError:
        future.cancel()
        log.warning(
            "resume_parser: LLM refinement exceeded %.2fs, falling back to heuristic",
            timeout,
        )
        return _with_parse_status(
            baseline,
            mode="basic",
            reason="timeout",
            started_at=started_at,
        )
    if not refined:
        return _with_parse_status(
            baseline,
            mode="basic",
            reason="llm_failed",
            started_at=started_at,
        )
    return _with_parse_status(
        _merge(baseline, refined),
        mode="ai_refined",
        reason="ai_completed",
        started_at=started_at,
    )
